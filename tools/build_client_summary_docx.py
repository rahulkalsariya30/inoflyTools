"""
tools/build_client_summary_docx.py

Build the short (~2 page) client-facing implementation summary as a Word file.

This is deliberately NOT the auditor package: it is the one-or-two page
"here is what we implemented against each requirement" note we hand to the
client. The full clause-by-clause submission stays in
Docs/audit/Inofly_Firmware_Security_Compliance_Document.docx.

Source of truth is Docs/CLIENT_IMPLEMENTATION_SUMMARY.md — edit that and rerun.

We reuse the Markdown->Word renderer from build_compliance_docx.py rather than
re-implementing it, so both deliverables format tables, bold and inline code the
same way. We skip that builder's cover page and table of contents on purpose:
at two pages they would cost more space than they add.

Usage:
    py -3 tools/build_client_summary_docx.py
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches

from build_compliance_docx import preprocess, render_markdown

REPO = Path(__file__).resolve().parents[1]
SOURCE = REPO / "Docs" / "CLIENT_IMPLEMENTATION_SUMMARY.md"
OUTPUT = REPO / "Docs" / "Inofly_Firmware_Security_Implementation_Summary.docx"


def main() -> None:
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(9.5)
    # Tight-but-readable body spacing; the default Word paragraph spacing alone
    # would push this past two pages.
    style.paragraph_format.space_after = Pt(4)

    # Narrow margins for the same reason — this is a summary sheet, not a report.
    for section in doc.sections:
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)

    # Word's stock Heading styles reserve a lot of space above each heading.
    # At this length that whitespace is what decides whether we land on two
    # pages or three, so tighten it rather than cutting content.
    for name, before, after, size in (
        ("Heading 1", 0, 8, 16),
        ("Heading 2", 10, 4, 12),
        ("Heading 3", 8, 3, 10.5),
    ):
        h = doc.styles[name]
        h.paragraph_format.space_before = Pt(before)
        h.paragraph_format.space_after = Pt(after)
        h.font.size = Pt(size)

    render_markdown(doc, preprocess(SOURCE.read_text(encoding="utf-8")))

    # The metadata table at the top is written as a headerless Markdown table
    # (`| | |`), which the shared renderer turns into an empty shaded header
    # band. Drop any such all-empty header row so it reads as a plain key/value
    # block instead of a table with a missing heading.
    for table in doc.tables:
        first = table.rows[0]
        if all(not c.text.strip() for c in first.cells):
            first._element.getparent().remove(first._element)

    # Table cells inherit the body paragraph spacing, which adds a blank line's
    # worth of padding to every row. Strip it back inside tables only.
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    para.paragraph_format.space_before = Pt(0)
                    para.paragraph_format.space_after = Pt(0)

    doc.save(str(OUTPUT))
    print(f"[OK] wrote {OUTPUT}  ({OUTPUT.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
