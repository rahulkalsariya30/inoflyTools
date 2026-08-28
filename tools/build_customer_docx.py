"""
tools/build_customer_docx.py

Build the customer-facing Firmware Security Compliance Document (Word).

This is the *customer* deliverable. It answers the same DGCA Section 7.1
conditions as the auditor package under Docs/audit/, and follows the section
structure and numbering of the reference compliance document supplied by the
client, so the two can be read side by side. The difference is depth, not
substance: no source paths, module names, requirement IDs or internal decision
records — a customer reads what the aircraft does and why it is trustworthy.

Source of truth is Docs/CUSTOMER_COMPLIANCE_DOCUMENT.md — edit that and rerun.

Figures: the Markdown carries `@@FIG <file.png>|<caption>` directive lines.
We render the Markdown in segments and drop the rendered PNG from
Docs/audit/diagrams/ in between, because the shared renderer has no image
syntax of its own (its diagrams are keyed on ASCII-art blocks, which this
document does not use). A missing PNG is skipped with a warning rather than
failing the build — regenerate them with tools/render_audit_diagrams.py.

Usage:
    py -3 tools/build_customer_docx.py
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches

from build_compliance_docx import preprocess, render_markdown, add_diagram

REPO = Path(__file__).resolve().parents[1]
SOURCE = REPO / "Docs" / "CUSTOMER_COMPLIANCE_DOCUMENT.md"
DIAGRAMS_DIR = REPO / "Docs" / "audit" / "diagrams"
OUTPUT = REPO / "Docs" / "Inofly_Firmware_Security_Compliance_Document_Customer.docx"

# Document-control block. These mirror the reference document's cover sheet;
# the names are left blank for wet signature, as on the reference.
DOC_TITLE = "Firmware Security Compliance Document"
DOC_NO = "12 - Software - INF"
DOC_ISSUE = "01"
REVISIONS = [("1", "26/08/2026", "Initial Release")]

FIG_DIRECTIVE = re.compile(r"^@@FIG\s+(\S+)\s*\|\s*(.+)$")


def _cell_text(cell, text, *, bold=False, size=10, center=True):
    para = cell.paragraphs[0]
    if center:
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)


def add_cover(doc):
    doc.add_paragraph()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Inofly")
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rs = sub.add_run("DGCA UAS Type Certification — Certification Scheme Section 7.1 (Level 1)")
    rs.font.size = Pt(11)
    rs.font.color.rgb = RGBColor(0x40, 0x40, 0x40)

    for _ in range(3):
        doc.add_paragraph()

    # Title / document number block
    head = doc.add_table(rows=2, cols=3)
    head.style = "Table Grid"
    head.alignment = WD_TABLE_ALIGNMENT.CENTER
    _cell_text(head.cell(0, 0), DOC_TITLE, bold=True, size=11)
    head.cell(0, 0).merge(head.cell(1, 0))
    _cell_text(head.cell(0, 1), "Doc No.:", bold=True)
    _cell_text(head.cell(0, 2), DOC_NO, bold=True)
    _cell_text(head.cell(1, 1), "Issue", bold=True)
    _cell_text(head.cell(1, 2), DOC_ISSUE, bold=True)

    doc.add_paragraph()

    # Version history
    hist = doc.add_table(rows=2 + len(REVISIONS), cols=3)
    hist.style = "Table Grid"
    hist.alignment = WD_TABLE_ALIGNMENT.CENTER
    _cell_text(hist.cell(0, 0), "Version History Table", bold=True, size=11)
    hist.cell(0, 0).merge(hist.cell(0, 2))
    for col, label in enumerate(("Revision No.:", "Revision Date:", "Description")):
        _cell_text(hist.cell(1, col), label, bold=True)
    for row, rev in enumerate(REVISIONS, start=2):
        for col, value in enumerate(rev):
            _cell_text(hist.cell(row, col), value, bold=True)

    doc.add_paragraph()

    # Sign-off block — third row left empty for the signature itself.
    sign = doc.add_table(rows=3, cols=2)
    sign.style = "Table Grid"
    sign.alignment = WD_TABLE_ALIGNMENT.CENTER
    _cell_text(sign.cell(0, 0), "Prepared by:", bold=True)
    _cell_text(sign.cell(0, 1), "Approved by:", bold=True)
    _cell_text(sign.cell(1, 0), "", bold=True)
    _cell_text(sign.cell(1, 1), "", bold=True)
    for cell in sign.rows[2].cells:
        cell.paragraphs[0].add_run("\n\n")

    doc.add_page_break()


def add_contents(doc):
    doc.add_heading("CONTENTS", level=1)
    p = doc.add_paragraph()
    run = p.add_run()
    fld_start = OxmlElement("w:fldChar")
    fld_start.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = r'TOC \o "1-3" \h \z \u'
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Right-click and choose “Update Field” to build the table of contents."
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    for el in (fld_start, instr, fld_sep, placeholder, fld_end):
        run._r.append(el)
    doc.add_page_break()


def render_body(doc, md: str):
    """Render the Markdown, splitting at @@FIG directives to insert figures."""
    buffer: list[str] = []
    for line in md.split("\n"):
        match = FIG_DIRECTIVE.match(line.strip())
        if not match:
            buffer.append(line)
            continue
        render_markdown(doc, "\n".join(buffer))
        buffer = []
        png, caption = DIAGRAMS_DIR / match.group(1), match.group(2)
        if png.exists():
            add_diagram(doc, png, caption)
        else:
            print(f"[warn] figure missing, skipped: {png.name}")
    if buffer:
        render_markdown(doc, "\n".join(buffer))


def main() -> None:
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    add_cover(doc)
    add_contents(doc)
    render_body(doc, preprocess(SOURCE.read_text(encoding="utf-8")))

    # The Section 1 summary table is 5 columns of very uneven text. Word's
    # autofit gives the two long columns too little room, so pin the widths and
    # repeat the header row on every page the table spans.
    for table in doc.tables:
        if len(table.columns) != 5:
            continue
        table.autofit = False
        for width, column in zip((0.5, 0.9, 1.6, 2.6, 0.9), table.columns):
            for cell in column.cells:
                cell.width = Inches(width)
        header = table.rows[0]._tr.get_or_add_trPr()
        repeat = OxmlElement("w:tblHeader")
        repeat.set(qn("w:val"), "true")
        header.append(repeat)

    # Table cells inherit body spacing; trimming it keeps the long §1 summary
    # table from running to twice the pages it needs.
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    para.paragraph_format.space_before = Pt(0)
                    para.paragraph_format.space_after = Pt(0)

    doc.save(str(OUTPUT))
    print(f"[OK] wrote {OUTPUT}  ({OUTPUT.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
